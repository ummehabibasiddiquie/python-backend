"""Sum a task's selected Excel column for range-based QA file/QC records."""

from __future__ import annotations

import csv
import io
from urllib.request import Request, urlopen

from openpyxl import load_workbook

from utils.qa_targets import QA_NEW_TARGETS_EFFECTIVE_FROM, _parse_work_date, parse_qa_target_ranges


def _norm(value) -> str:
    return " ".join(str(value or "").strip().lower().split())


def _loose(value) -> str:
    return " ".join(part for part in _norm(value).split() if part != "add")


def _find_column(headers, column_name: str) -> int | None:
    want = _norm(column_name)
    want_loose = _loose(column_name)
    loose_idx = None
    for idx, header in enumerate(headers):
        if _norm(header) == want:
            return idx
        if loose_idx is None and want_loose and _loose(header) == want_loose:
            loose_idx = idx
    return loose_idx


def _cell_number(value) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).replace(",", "").strip()
    try:
        return float(text)
    except ValueError:
        return 0.0


def _row_has_value(row) -> bool:
    for cell in row:
        if cell is not None and str(cell).strip() != "":
            return True
    return False


def sum_workbook_column(content: bytes, column_name: str, ext: str) -> float | None:
    ext = (ext or "").lower()
    if ext == ".csv":
        text = content.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
        if not rows:
            return None
        col = _find_column(rows[0], column_name)
        if col is None:
            return None
        total = 0.0
        for row in rows[1:]:
            if not _row_has_value(row):
                continue
            total += _cell_number(row[col] if col < len(row) else None)
        return total

    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    try:
        sheet = workbook.worksheets[0]
        row_iter = sheet.iter_rows(values_only=True)
        headers = next(row_iter, None)
        if not headers:
            return None
        col = _find_column(headers, column_name)
        if col is None:
            return None
        total = 0.0
        for row in row_iter:
            if not _row_has_value(row):
                continue
            total += _cell_number(row[col] if col < len(row) else None)
        return total
    finally:
        workbook.close()


def _download(url: str) -> bytes | None:
    if not url or not str(url).startswith("http"):
        return None
    req = Request(url, headers={"User-Agent": "HRMS-QA-Report"})
    with urlopen(req, timeout=25) as resp:
        return resp.read()


def ensure_qc_object_count_column(cursor) -> None:
    cursor.execute("SHOW COLUMNS FROM qc_records LIKE 'qc_object_count'")
    if cursor.fetchone():
        return
    cursor.execute(
        """
        ALTER TABLE qc_records
        ADD COLUMN qc_object_count DECIMAL(14,2) NULL AFTER object_count
        """
    )


def fill_missing_range_counts(cursor, rec: dict, work_date) -> None:
    """
    If this range task was saved before the column sum existed, read the
    agent file once and store the selected-column total.
    """
    wd = _parse_work_date(work_date)
    if wd is not None and wd < QA_NEW_TARGETS_EFFECTIVE_FROM:
        return
    if rec.get("object_count") is not None:
        return
    column = str(rec.get("qa_count_column") or "").strip()
    if not column or not parse_qa_target_ranges(rec.get("qa_target_ranges")):
        return
    url = rec.get("whole_file_path") or ""
    ext = ""
    path = url.split("?", 1)[0]
    if "." in path:
        ext = "." + path.rsplit(".", 1)[-1].lower()
    if ext not in (".xlsx", ".csv"):
        return
    try:
        content = _download(url)
        if not content:
            return
        total = sum_workbook_column(content, column, ext)
    except Exception:
        return
    if total is None:
        return

    file_rows = float(rec.get("file_record_count") or 0)
    qc_rows = float(rec.get("qc_generated_count") or 0)
    if rec.get("qc_object_count") is not None:
        qc_sum = rec.get("qc_object_count")
    elif file_rows <= 0 or qc_rows >= file_rows:
        qc_sum = total
    else:
        qc_sum = total * qc_rows / file_rows

    rec["object_count"] = total
    rec["qc_object_count"] = qc_sum
    qc_id = rec.get("qc_record_id")
    if not qc_id:
        return
    cursor.execute(
        """
        UPDATE qc_records
        SET object_count=%s, qc_object_count=%s
        WHERE id=%s AND object_count IS NULL
        """,
        (round(float(total), 2), round(float(qc_sum), 2), int(qc_id)),
    )
