"""Excel export of the monthly KRA. Calculated cells are formulas, not pasted values."""

from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from utils.kra import (
    PRODUCTIVITY_HOURS,
    QUALITY_MIN,
    REPORTING_FULL_MAX,
    REPORTING_FULL_SCORE,
    REPORTING_PARTIAL_MAX,
    REPORTING_PARTIAL_SCORE,
    TRACKER_MIN,
    WEIGHT_PRODUCTIVITY,
    WEIGHT_QUALITY,
    WEIGHT_REPORTING,
    WEIGHT_SCHEDULE,
    WEIGHT_TIMELINESS,
)

ATTENDANCE_FILL = {
    "PRESENT": "D1FAE5",
    "HALF DAY": "FEF3C7",
    "ABSENT": "FEE2E2",
    "LEAVE": "FEF9C3",
    "WEEK OFF": "E0F2FE",
    "WFH": "CCFBF1",
    "UNROSTERED": "FFEDD5",
    "HOLIDAY": "EDE9FE",
}
ATTENDANCE_FONT = {
    "PRESENT": "065F46",
    "HALF DAY": "92400E",
    "ABSENT": "991B1B",
    "LEAVE": "854D0E",
    "WEEK OFF": "075985",
    "WFH": "115E59",
    "UNROSTERED": "9A3412",
    "HOLIDAY": "5B21B6",
}
ATTENDANCE_ROW_TINT = {
    "HALF DAY": "FFFBEB",
    "ABSENT": "FEF2F2",
    "LEAVE": "FEFCE8",
    "WEEK OFF": "F0F9FF",
    "WFH": "F0FDFA",
    "UNROSTERED": "FFF7ED",
    "HOLIDAY": "F5F3FF",
}

NAVY = "1B3A4B"
TEAL = "0F6E6B"
SLATE = "F4F7F8"
AMBER = "FFF6DE"
LINE = "D5DEE3"
GREEN = "E5F6ED"
RED = "FDECEC"
MUTED = "5C6B73"
WHITE = "FFFFFF"

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
wrap = Alignment(wrap_text=True, vertical="center")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(size=11, bold=False, color="1F2933"):
    return Font(name="Calibri", size=size, bold=bold, color=color)


def _header(cell):
    cell.fill = _fill(NAVY)
    cell.font = _font(10, True, WHITE)
    cell.alignment = center
    cell.border = thin


def _safe_name(value: str) -> str:
    cleaned = re.sub(r"[^\w\s-]", "", value or "Agent").strip()
    return re.sub(r"\s+", " ", cleaned) or "Agent"


def build_kra_workbook(report: dict) -> tuple[io.BytesIO, str]:
    import logging
    logger = logging.getLogger(__name__)
    
    days = report.get("days") or []
    first = 6
    last = first + max(len(days), 1) - 1

    logger.info(f"KRA Excel - Building workbook for {report.get('user_name')}, days: {len(days)}, first: {first}, last: {last}")

    wb = Workbook()
    rules = wb.active
    rules.title = "Rules"
    daily = wb.create_sheet("Daily Log")
    score = wb.create_sheet("KRA Score")
    _write_rules(rules)
    totals_rows = _write_daily(daily, report, first, last)
    _write_score(score, report, totals_rows)

    # Verify formulas were written
    logger.info(f"KRA Excel - Daily Log sheet rows: {daily.max_row}, Score sheet rows: {score.max_row}")
    logger.info(f"KRA Excel - Totals rows: {totals_rows}")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"KRA {_safe_name(report.get('user_name'))} - {report.get('month_year')}"
    if report.get("period_end"):
        filename += f" through {report.get('period_end')}"
    filename += ".xlsx"
    
    logger.info(f"KRA Excel - Workbook saved, filename: {filename}")
    return output, filename


def _write_rules(ws):
    ws.sheet_view.showGridLines = False
    ws["A1"] = "KRA calculation rules"
    ws["A1"].font = _font(18, True, NAVY)
    ws.merge_cells("A1:C1")
    ws["A2"] = "Daily Log and KRA Score read these cells. Change a threshold here and the scores recalculate."
    ws["A2"].font = _font(11, False, MUTED)
    ws.merge_cells("A2:C2")

    for col, label in enumerate(("Setting", "Value", "What it does"), 1):
        _header(ws.cell(4, col, label))

    rows = [
        ("Productivity hours", PRODUCTIVITY_HOURS, "Billable hours at or above this = YES. Below this = NO."),
        ("Quality minimum", QUALITY_MIN, "QC score at or above this = YES. Below this = NO. Sheet rule: equals or exceeds 98%."),
        ("Minimum trackers", TRACKER_MIN, "A Present, Half Day, WFH, or Unrostered day with fewer trackers is one non-compliance."),
        ("Full reporting band up to", REPORTING_FULL_MAX, f"0 to this many instances still earns {REPORTING_FULL_SCORE}%."),
        ("Partial reporting band up to", REPORTING_PARTIAL_MAX, f"Above the full band, up to this number, earns {REPORTING_PARTIAL_SCORE}%. More than this earns 0%."),
        ("Full reporting score", REPORTING_FULL_SCORE, "Points earned inside the first band."),
        ("Partial reporting score", REPORTING_PARTIAL_SCORE, "Points earned inside the second band."),
        ("Productivity weight", WEIGHT_PRODUCTIVITY, "(YES days / working days) x this weight."),
        ("Quality weight", WEIGHT_QUALITY, "(quality YES days / working days) x this weight."),
        ("Schedule weight", WEIGHT_SCHEDULE, "(present days / working days) x this weight."),
        ("Reporting weight", WEIGHT_REPORTING, "Band score. The maximum is this weight."),
        ("Timeliness weight", WEIGHT_TIMELINESS, "Always included in total weight (100%). Earned stays blank until typed in KRA Score!D13 (0 to 10); then total earned updates."),
    ]
    for idx, (label, value, note) in enumerate(rows, 5):
        ws.cell(idx, 1, label).font = _font(11, True)
        value_cell = ws.cell(idx, 2, value)
        value_cell.font = _font(12, True, TEAL)
        value_cell.alignment = center
        value_cell.fill = _fill(AMBER)
        ws.cell(idx, 3, note).font = _font(11)
        for col in range(1, 4):
            ws.cell(idx, col).border = thin
            ws.cell(idx, col).alignment = wrap

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 88
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:4"
    ws.page_setup.horizontalCentered = True


def _write_daily(ws, report, first, last):
    ws.sheet_view.showGridLines = True
    ws.merge_cells("A1:I1")
    ws["A1"] = f"KRA daily log — {report.get('user_name') or 'Agent'}"
    ws["A1"].font = _font(18, True, NAVY)
    team = report.get("team_name") or "—"
    ws.merge_cells("A2:I2")
    period = report.get("period_end") or ""
    as_of = f"   ·   Through {period}" if period else ""
    ws["A2"] = f"{report.get('month_year')}   ·   Team {team}{as_of}   ·   Generated from HRMS. Yellow cells are formulas."
    ws["A2"].font = _font(11, False, MUTED)

    headers = [
        "Date", "Day", "Attendance", "Billable Hours", "Productivity Achieved",
        "Quality Score", "Quality Achieved", "Trackers", "Note",
    ]
    for col, label in enumerate(headers, 1):
        _header(ws.cell(5, col, label))

    for offset, row in enumerate(report.get("days") or []):
        r = first + offset
        ws.cell(r, 1, row.get("work_date"))
        ws.cell(r, 2, row.get("day"))
        ws.cell(r, 3, row.get("attendance") or None)
        ws.cell(r, 4, row.get("billable_hours"))
        ws.cell(
            r, 5,
            f'=IF(OR(C{r}="PRESENT",C{r}="HALF DAY",C{r}="WFH",C{r}="ABSENT",C{r}="UNROSTERED",AND(ISNUMBER(D{r}),D{r}>0)),IF(N(D{r})>=Rules!$B$5,"YES","NO"),"")',
        )
        ws.cell(r, 6, row.get("qc_score"))
        ws.cell(r, 7, f'=IF(F{r}="","",IF(F{r}>=Rules!$B$6,"YES","NO"))')
        trackers = row.get("tracker_count")
        if trackers is None and (row.get("attendance") or "") in ("PRESENT", "HALF DAY", "WFH", "UNROSTERED"):
            trackers = 0
        ws.cell(r, 8, trackers)
        ws.cell(r, 9, row.get("note") or None)
        for col in range(1, 10):
            cell = ws.cell(r, col)
            cell.border = thin
            cell.font = _font(10)
            cell.alignment = center if col < 9 else Alignment(vertical="center", wrap_text=True)
        ws.cell(r, 4).number_format = "0.00"
        ws.cell(r, 6).number_format = "0.00"
        ws.cell(r, 5).fill = _fill(AMBER)
        ws.cell(r, 7).fill = _fill(AMBER)
        ws.cell(r, 9).fill = _fill("F8FBFC")
        status = (row.get("attendance") or "").strip().upper()
        if status in ATTENDANCE_FILL:
            ws.cell(r, 3).fill = _fill(ATTENDANCE_FILL[status])
            ws.cell(r, 3).font = _font(10, True, ATTENDANCE_FONT.get(status, "1F2933"))
            row_tint = ATTENDANCE_ROW_TINT.get(status)
            if row_tint:
                for col in (1, 2, 4, 6, 8):
                    ws.cell(r, col).fill = _fill(row_tint)
        elif offset % 2 == 1:
            for col in (1, 2, 4, 6, 8):
                ws.cell(r, col).fill = _fill(SLATE)

    span = f"$C${first}:$C${last}"
    prod_col = f"$E${first}:$E${last}"
    qual_col = f"$G${first}:$G${last}"
    tracker_col = f"$H${first}:$H${last}"
    hours_col = f"$D${first}:$D${last}"
    score_col = f"$F${first}:$F${last}"

    def status_count(status):
        return f'COUNTIF({span},"{status}")'

    working = "+".join(status_count(s) for s in ("PRESENT", "HALF DAY", "ABSENT", "WFH", "UNROSTERED"))
    present = "+".join(status_count(s) for s in ("PRESENT", "HALF DAY", "WFH"))
    low_tracker = "+".join(
        f'COUNTIFS({span},"{status}",{tracker_col},"<"&Rules!$B$7)'
        for status in ("PRESENT", "HALF DAY", "WFH", "UNROSTERED")
    )

    # End-of-listing totals — same labels as the shared sheet, laid out as
    # three compact label|value pairs (no blank columns between label and value).
    avg_row = last + 2
    counts_row = last + 3
    working_row = last + 4
    score_row = last + 5
    low_row = last + 6
    totals_rows = {
        "avg_row": avg_row,
        "counts_row": counts_row,
        "working_row": working_row,
        "score_row": score_row,
        "productivity_yes": counts_row,
        "quality_yes": counts_row,
        "present_days": counts_row,
        "working_days": working_row,
        "low_tracker_days": low_row,
        "productivity_score": score_row,
        "quality_score": score_row,
        "schedule_score": score_row,
    }

    def style_label(cell, bold=True):
        cell.font = _font(10, bold, NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="left")
        cell.border = thin
        cell.fill = _fill(SLATE)

    def style_value(cell):
        cell.font = _font(12, True, TEAL)
        cell.fill = _fill(AMBER)
        cell.alignment = center
        cell.border = thin

    def merge_label(row, c1, c2, text):
        if c2 > c1:
            ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
        cell = ws.cell(row, c1, text)
        style_label(cell)
        for col in range(c1, c2 + 1):
            ws.cell(row, col).border = thin
            ws.cell(row, col).fill = _fill(SLATE)

    def put_value(row, col, value, number_format=None):
        cell = ws.cell(row, col, value)
        style_value(cell)
        if number_format:
            cell.number_format = number_format
        return cell

    # Total/Average — values sit under the matching daily columns
    # D = Billable Hours, F = Quality Score
    merge_label(avg_row, 1, 2, "Total/Average")
    put_value(avg_row, 4, f"=SUM({hours_col})", "0.00")
    put_value(avg_row, 6, f'=IF(COUNT({score_col})=0,"",AVERAGE({score_col}))', "0.00")

    # Achieved counts — Productivity | Quality | Present
    merge_label(counts_row, 1, 2, "No. of Days Productivity Achieved")
    put_value(counts_row, 3, f'=COUNTIF({prod_col},"YES")')
    merge_label(counts_row, 4, 5, "No. of Days Quality Achieved")
    put_value(counts_row, 6, f'=COUNTIF({qual_col},"YES")')
    merge_label(counts_row, 7, 8, "Present Days")
    put_value(counts_row, 9, f"={present}")

    # Working days under each group
    merge_label(working_row, 1, 2, "Total Working Days")
    put_value(working_row, 3, f"={working}")
    merge_label(working_row, 4, 5, "Total Working Days")
    put_value(working_row, 6, f"={working}")
    merge_label(working_row, 7, 8, "Total Working Days")
    put_value(working_row, 9, f"={working}")

    # Scores
    merge_label(score_row, 1, 2, "Score")
    put_value(
        score_row, 3,
        f'=IF(C{working_row}=0,"",C{counts_row}/C{working_row}*Rules!$B$12)',
        "0.00",
    )
    merge_label(score_row, 4, 5, "Score")
    put_value(
        score_row, 6,
        f'=IF(F{working_row}=0,"",F{counts_row}/F{working_row}*Rules!$B$13)',
        "0.00",
    )
    merge_label(score_row, 7, 8, "Score")
    put_value(
        score_row, 9,
        f'=IF(I{working_row}=0,"",I{counts_row}/I{working_row}*Rules!$B$14)',
        "0.00",
    )

    merge_label(low_row, 1, 2, "Days under 7 trackers")
    put_value(low_row, 3, f"={low_tracker}")

    for r in (avg_row, counts_row, working_row, score_row, low_row):
        ws.row_dimensions[r].height = 24

    # Right-side attendance counts (beside the daily log, after Note)
    _header(ws.cell(5, 11, "Attendance"))
    _header(ws.cell(5, 12, "Count"))
    for idx, status in enumerate(("PRESENT", "HALF DAY", "ABSENT", "LEAVE", "WEEK OFF", "WFH", "UNROSTERED", "HOLIDAY")):
        r = first + idx
        label = ws.cell(r, 11, status)
        label.font = _font(10, True, ATTENDANCE_FONT.get(status, "1F2933"))
        label.border = thin
        label.fill = _fill(ATTENDANCE_FILL.get(status, WHITE))
        label.alignment = Alignment(vertical="center")
        cell = ws.cell(r, 12, f"={status_count(status)}")
        cell.font = _font(11, True, ATTENDANCE_FONT.get(status, "1F2933"))
        cell.alignment = center
        cell.border = thin
        cell.fill = _fill(ATTENDANCE_FILL.get(status, AMBER))

    _header(ws.cell(first + 9, 11, "Productivity"))
    _header(ws.cell(first + 9, 12, "Count"))
    yes_r, no_r = first + 10, first + 11
    ws.cell(yes_r, 11, "YES").font = _font(10, True, "166534")
    ws.cell(yes_r, 11).fill = _fill(GREEN)
    ws.cell(yes_r, 11).border = thin
    put_value(yes_r, 12, f'=COUNTIF({prod_col},"YES")')
    ws.cell(no_r, 11, "NO").font = _font(10, True, "991B1B")
    ws.cell(no_r, 11).fill = _fill(RED)
    ws.cell(no_r, 11).border = thin
    put_value(no_r, 12, f'=COUNTIF({prod_col},"NO")')

    for status, color in ATTENDANCE_FILL.items():
        ws.conditional_formatting.add(
            f"C{first}:C{last}",
            FormulaRule(
                formula=[f'C{first}="{status}"'],
                fill=_fill(color),
                font=_font(10, True, ATTENDANCE_FONT[status]),
            ),
        )

    yes_rule = FormulaRule(formula=[f'E{first}="YES"'], fill=_fill(GREEN), font=_font(10, True, "166534"))
    no_rule = FormulaRule(formula=[f'E{first}="NO"'], fill=_fill(RED), font=_font(10, True, "991B1B"))
    ws.conditional_formatting.add(f"E{first}:E{last}", yes_rule)
    ws.conditional_formatting.add(f"E{first}:E{last}", no_rule)
    ws.conditional_formatting.add(f"G{first}:G{last}", FormulaRule(formula=[f'G{first}="YES"'], fill=_fill(GREEN), font=_font(10, True, "166534")))
    ws.conditional_formatting.add(f"G{first}:G{last}", FormulaRule(formula=[f'G{first}="NO"'], fill=_fill(RED), font=_font(10, True, "991B1B")))
    ws.conditional_formatting.add(
        f"H{first}:H{last}",
        FormulaRule(formula=[f'AND(ISNUMBER(H{first}),H{first}<Rules!$B$7)'], fill=_fill(RED), font=_font(10, True, "991B1B")),
    )

    # Daily columns stay compact; K:L are the right-side counts only
    for col, width in {
        "A": 12, "B": 11, "C": 12, "D": 13, "E": 18,
        "F": 12, "G": 14, "H": 10, "I": 22, "J": 2,
        "K": 14, "L": 9,
    }.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:I{last}"
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[5].height = 32
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:5"
    ws.page_setup.horizontalCentered = True
    ws.oddFooter.right.text = "Page &P of &N"
    return totals_rows


def _write_score(ws, report, totals_rows):
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F1")
    ws["A1"] = f"KRA score — {report.get('user_name') or 'Agent'} — {report.get('month_year')}"
    ws["A1"].font = _font(18, True, NAVY)
    ws.merge_cells("A2:F2")
    ws["A2"] = "Points 1–4 come from HRMS. Point 5 (Timeliness) weight is always in the 100% total; leave D13 blank until you type the earned score (0–10)."
    ws["A2"].font = _font(11, False, MUTED)

    for col, label in enumerate(("Sr", "Objective", "Weight %", "Earned %", "How it is calculated", "% of weight"), 1):
        _header(ws.cell(4, col, label))

    counts_row = totals_rows["counts_row"]
    working_row = totals_rows["working_row"]
    score_row = totals_rows["score_row"]
    low_row = totals_rows["low_tracker_days"]
    prod_score = f"'Daily Log'!C{score_row}"
    qual_score = f"'Daily Log'!F{score_row}"
    sched_score = f"'Daily Log'!I{score_row}"
    low_tracker = f"'Daily Log'!C{low_row}"

    rows = [
        (5, 1, "Meeting Daily Productivity", "=Rules!$B$12",
         f"={prod_score}",
         f"(No. of Days Productivity Achieved / Total Working Days) x 33  — Daily Log!C{counts_row}/C{working_row}"),
        (7, 2, "Delivering Right Quality Everyday", "=Rules!$B$13",
         f"={qual_score}",
         f"(No. of Days Quality Achieved / Total Working Days) x 33  — Daily Log!F{counts_row}/F{working_row}"),
        (9, 3, "Schedule Adherence - Rostered Attendance", "=Rules!$B$14",
         f"={sched_score}",
         f"(Present Days / Total Working Days) x 10  — Daily Log!I{counts_row}/I{working_row}"),
        (11, 4, "Adherence to Reporting in TimeChamp, Project Tracker, and Keka", "=Rules!$B$15",
         '=IF(I14<=Rules!$B$8,Rules!$B$10,IF(I14<=Rules!$B$9,Rules!$B$11,0))',
         "0-3 instances = 14%. 4-6 = 7%. More than 6 = 0%. Verbal warning = 1, email = 2, letter = 3."),
        (13, 5, "Timeliness - Adherence to break and login schedule", "=Rules!$B$16",
         None,
         "Weight 10% always counts in total. Type earned in D13 (0–10) when ready — total earned then updates. Example: up to 3 non-compliances = 10%, 4–5 = 5%, more than 5 = 0%."),
    ]
    for r, sr, title, weight, earned, how in rows:
        ws.cell(r, 1, sr)
        ws.cell(r, 2, title)
        ws.cell(r, 3, weight)
        if earned:
            ws.cell(r, 4, earned)
        ws.cell(r, 5, how)
        ws.cell(r, 6, f'=IF(OR(C{r}="",C{r}=0,D{r}=""),"",D{r}/C{r})')
        for col in range(1, 7):
            cell = ws.cell(r, col)
            cell.border = thin
            cell.font = _font(11, col in (1, 4))
            cell.alignment = center if col not in (2, 5) else wrap
        ws.cell(r, 3).fill = _fill(SLATE)
        ws.cell(r, 4).fill = _fill(AMBER if earned else "FFF2CC")
        ws.cell(r, 6).fill = _fill(AMBER)
        ws.cell(r, 3).number_format = "0"
        ws.cell(r, 4).number_format = "0.00"
        ws.cell(r, 6).number_format = "0.0%"
        ws.row_dimensions[r].height = 36

    ws["D13"] = None
    ws["D13"].fill = _fill("FFF2CC")
    ws["D13"].number_format = "0.00"
    ws["D13"].border = thin
    ws["D13"].alignment = center
    timeliness_dv = DataValidation(
        type="decimal",
        operator="between",
        formula1="0",
        formula2="10",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Timeliness score",
        error="Enter a number from 0 to 10, or leave blank.",
    )
    timeliness_dv.add("D13")
    ws.add_data_validation(timeliness_dv)

    ws["H10"] = "Low tracker days"
    ws["I10"] = f"={low_tracker}"
    ws["H11"] = "Verbal warnings (x1)"
    ws["I11"] = 0
    ws["H12"] = "Email warnings (x2)"
    ws["I12"] = 0
    ws["H13"] = "Letter warnings (x3)"
    ws["I13"] = 0
    ws["H14"] = "Total instances"
    ws["I14"] = "=I10+I11+I12*2+I13*3"
    for r in range(10, 15):
        ws.cell(r, 8).font = _font(10, r == 14)
        ws.cell(r, 8).border = thin
        ws.cell(r, 8).alignment = Alignment(vertical="center")
        ws.cell(r, 9).font = _font(11, True, TEAL)
        ws.cell(r, 9).alignment = center
        ws.cell(r, 9).border = thin
        ws.cell(r, 9).fill = _fill(AMBER if r in (10, 14) else "FFF2CC")

    ws["B15"] = "Total earned"
    ws["C15"] = "=C5+C7+C9+C11+C13"
    ws["D15"] = '=IF(D5="",0,D5)+IF(D7="",0,D7)+IF(D9="",0,D9)+D11+IF(D13="",0,D13)'
    ws["E15"] = "Weight always includes Timeliness 10% (100% total). Leave D13 blank until you type earned; then total earned updates."
    ws["F15"] = '=IF(C15=0,"",D15/C15)'
    for col in range(2, 7):
        cell = ws.cell(15, col)
        cell.font = _font(12, True, WHITE)
        cell.fill = _fill(TEAL)
        cell.alignment = center if col != 5 else wrap
        cell.border = thin
    ws["C15"].number_format = "0"
    ws["D15"].number_format = "0.00"
    ws["F15"].number_format = "0.0%"
    ws.row_dimensions[15].height = 28

    ws.merge_cells("A17:F17")
    ws["A17"] = "User input cells (light yellow): Timeliness earned % in D13, and warning counts in I11–I13. Other yellow cells are formulas."
    ws["A17"].font = _font(10, False, MUTED)

    for col, width in {"A": 8, "B": 64, "C": 12, "D": 14, "E": 78, "F": 14, "G": 3, "H": 26, "I": 18}.items():
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:4"
    ws.page_setup.horizontalCentered = True
    ws.oddFooter.right.text = "Page &P of &N"
