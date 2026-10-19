"""Rebuild qc_score and error_list for QC records that a correction save wiped.

Before qc-backend commit 4e88694 (12 Oct 2026), marking a correction completed
wrote qc_score = NULL and error_list = NULL on qc_records. The marked QC sample
saved at first QC still lists the errors on each row, so the score can be
recalculated the same way the QC form does it.

Usage (from backend-api, uses the DB_* settings in .env):
    python scripts/repair_qc_null_scores.py                 # dry run
    python scripts/repair_qc_null_scores.py --ids 6966      # dry run, one record
    python scripts/repair_qc_null_scores.py --apply         # write changes
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import os
import re
import sys
import urllib.request
from datetime import datetime

import openpyxl
from dotenv import load_dotenv

load_dotenv()
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from scripts.hostinger_db import connect  # noqa: E402


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Recalculate wiped QC scores from the saved QC sample files")
    p.add_argument("--host", default=os.getenv("DB_HOST", "localhost"))
    p.add_argument("--port", type=int, default=int(os.getenv("DB_PORT", "3306")))
    p.add_argument("--user", default=os.getenv("DB_USERNAME", "root"))
    p.add_argument("--database", default=os.getenv("DB_DATABASE", "tfs_hrms"))
    p.add_argument("--password", default=None)
    p.add_argument("--ids", default="", help="Comma-separated qc_records ids. Default: every record with qc_score NULL")
    p.add_argument("--apply", action="store_true", help="Write the recalculated score and error list")
    p.add_argument("--report", default="", help="CSV report path (default: scripts/repair_qc_null_scores_<time>.csv)")
    args = p.parse_args()
    if args.password is None and args.host == os.getenv("DB_HOST", "localhost"):
        args.password = os.getenv("DB_PASSWORD")
    return args


def cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def norm(text) -> str:
    return " ".join(cell_text(text).replace("_", " ").split()).lower()


def is_qc_code_header(name) -> bool:
    n = norm(name)
    compact = "".join(ch for ch in n if ch.isalnum())
    return n == "qc code" or compact == "qccode" or ("qc" in n and "code" in n)


def download_workbook(url: str):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    data = urllib.request.urlopen(req, timeout=90).read()
    return openpyxl.load_workbook(io.BytesIO(data), data_only=True)


def load_afd(cursor, project_id):
    cursor.execute(
        "SELECT pc.afd_id FROM project p "
        "JOIN project_category pc ON pc.project_category_id = p.project_category_id "
        "WHERE p.project_id = %s",
        (project_id,),
    )
    row = cursor.fetchone()
    if not row or row["afd_id"] is None:
        return None
    cursor.execute(
        "SELECT qc_afd_id, afd_name, afd_points, afd_category_id FROM qc_afd WHERE afd_id = %s",
        (row["afd_id"],),
    )
    items = cursor.fetchall()
    categories = [i for i in items if int(i["afd_category_id"] or 0) == 0]
    by_cat = {}
    for cat in categories:
        subs = [i for i in items if int(i["afd_category_id"] or 0) == int(cat["qc_afd_id"])]
        by_cat[norm(cat["afd_name"])] = {
            "name": cat["afd_name"],
            "subs": {norm(s["afd_name"]): {"name": s["afd_name"], "points": float(s["afd_points"] or 0)} for s in subs},
        }
    return by_cat


def resolve_error(afd, category, subcategory, stored_points):
    """Return (category_key, category_name, subcategory_name, points)."""
    cat_key = norm(category)
    sub_key = norm(subcategory)
    if afd:
        cat = afd.get(cat_key)
        if cat and sub_key in cat["subs"]:
            sub = cat["subs"][sub_key]
            return cat_key, cat["name"], sub["name"], sub["points"]
        for key, cat in afd.items():
            if sub_key in cat["subs"]:
                sub = cat["subs"][sub_key]
                return key, cat["name"], sub["name"], sub["points"]
    return cat_key, category, subcategory, float(stored_points or 0)


def record_score(row_errors, afd) -> float:
    """Same rules as QCFormPage calculateRecordScore."""
    if not afd:
        return 100.0
    if any(err["points"] >= 100 for err in row_errors):
        return 0.0
    scores = []
    for cat_key in afd:
        deduction = sum(err["points"] for err in row_errors if err["cat_key"] == cat_key)
        scores.append(max(0.0, 100.0 - deduction))
    return sum(scores) / len(scores) if scores else 100.0


def parse_label(label: str):
    text = cell_text(label)
    if " - " in text:
        category, subcategory = text.split(" - ", 1)
        return category.strip(), subcategory.strip()
    return "", text


class NoErrorColumn(Exception):
    """Sample saved before errors were written into the file."""


SUMMARY_RE = re.compile(r"^(.*)\((\d+)\)\s*$")


def find_error_column(ws, headers):
    lower = [h.lower() for h in headers]
    if "errors" in lower:
        return lower.index("errors")
    for row in ws.iter_rows(min_row=2):
        for idx, cell in enumerate(row):
            if norm(cell.value) == "error list":
                return idx
    return None


def read_summary(ws, errors_col):
    """Counts from the 'Error List' block under the errors column."""
    counts = []
    started = False
    for r in range(2, ws.max_row + 1):
        value = cell_text(ws.cell(r, errors_col + 1).value)
        if not started:
            started = norm(value) == "error list"
            continue
        if not value:
            break
        if norm(value) == "no errors":
            return []
        match = SUMMARY_RE.match(value)
        if match:
            counts.append((match.group(1).strip(), int(match.group(2))))
        else:
            counts.append((value, 1))
    return counts if started else None


def read_sample(wb, afd):
    """Return (sample_row_count, error_list, per_row_errors, summary_counts)."""
    ws = wb.worksheets[0]
    headers = [cell_text(c.value) for c in ws[1]]
    lower = [h.lower() for h in headers]
    qc_errors_col = lower.index("qc_errors") if "qc_errors" in lower else None
    errors_col = find_error_column(ws, headers)
    if errors_col is None and qc_errors_col is None:
        raise NoErrorColumn("sample file has no error column")

    excel_row_col = lower.index("_excel_row") if "_excel_row" in lower else None
    qc_code_col = next((i for i, h in enumerate(headers) if is_qc_code_header(h)), None)
    skip_cols = {i for i in (qc_errors_col, errors_col, excel_row_col) if i is not None}
    skip_cols |= {i for i, h in enumerate(lower) if h == "has_error"}

    # Older annotations wrote each label one row too high, so the header cell
    # holds the first record's error instead of "Errors".
    shifted = (
        errors_col is not None
        and qc_errors_col is None
        and headers[errors_col] != ""
        and norm(headers[errors_col]) != "errors"
    )

    def errors_cell(sheet_row):
        if errors_col is None:
            return ""
        source_row = sheet_row - 1 if shifted else sheet_row
        return cell_text(ws.cell(source_row, errors_col + 1).value)

    error_list = []
    per_row = []
    for r in range(2, ws.max_row + 1):
        values = [c.value for c in ws[r]]
        if errors_col is not None and errors_col < len(values) and norm(values[errors_col]) == "error list":
            break
        data_values = [v for i, v in enumerate(values) if i not in skip_cols and cell_text(v) != ""]
        if not data_values:
            break

        qc_code = cell_text(values[qc_code_col]) if qc_code_col is not None and qc_code_col < len(values) else ""
        excel_row = r
        if excel_row_col is not None and excel_row_col < len(values):
            try:
                excel_row = int(float(values[excel_row_col]))
            except (TypeError, ValueError):
                pass

        raw_errors = []
        if qc_errors_col is not None and qc_errors_col < len(values) and cell_text(values[qc_errors_col]):
            try:
                parsed = json.loads(cell_text(values[qc_errors_col]))
                if isinstance(parsed, list):
                    raw_errors = [
                        (e.get("category", ""), e.get("subcategory", ""), e.get("points"))
                        for e in parsed
                        if isinstance(e, dict)
                    ]
            except json.JSONDecodeError:
                raw_errors = []
        if not raw_errors:
            for label in errors_cell(r).split(";"):
                if label.strip() and norm(label) not in ("error list", "no errors"):
                    category, subcategory = parse_label(label)
                    raw_errors.append((category, subcategory, None))

        row_errors = []
        for category, subcategory, stored_points in raw_errors:
            cat_key, cat_name, sub_name, points = resolve_error(afd, category, subcategory, stored_points)
            row_errors.append({"cat_key": cat_key, "points": points})
            entry = {"row": excel_row}
            if qc_code:
                entry["qc_code"] = qc_code
            entry.update(
                {
                    "category": cat_name,
                    "subcategory": sub_name,
                    "error": f"{cat_name} - {sub_name}",
                    "points": points if points != int(points) else int(points),
                }
            )
            error_list.append(entry)
        per_row.append(row_errors)

    summary = read_summary(ws, errors_col) if errors_col is not None else None
    return len(per_row), error_list, per_row, summary


def score_from_summary(summary, row_count, afd):
    """Score from the Error List counts. Same result as per-row scoring unless a record hits 0."""
    if not afd:
        return 100.0, False
    cat_count = len(afd) or 1
    total = 0.0
    fatal = False
    for label, count in summary:
        category, subcategory = parse_label(label)
        _, _, _, points = resolve_error(afd, category, subcategory, None)
        if points >= 100:
            fatal = True
        total += points * count
    return 100.0 - total / (cat_count * row_count), fatal


def sample_urls(cursor, rec):
    urls = []
    if rec.get("qc_file_path"):
        urls.append(rec["qc_file_path"])
    cursor.execute(
        "SELECT qc_file_path FROM qc_correction_history WHERE qc_record_id = %s "
        "AND qc_file_path IS NOT NULL AND qc_file_path <> '' ORDER BY correction_count ASC",
        (rec["id"],),
    )
    urls += [r["qc_file_path"] for r in cursor.fetchall()]
    cursor.execute(
        "SELECT rework_sample_file FROM qc_rework_history WHERE qc_record_id = %s "
        "AND rework_sample_file IS NOT NULL AND rework_sample_file <> '' ORDER BY rework_count ASC",
        (rec["id"],),
    )
    urls += [r["rework_sample_file"] for r in cursor.fetchall()]
    seen = []
    for url in urls:
        if url not in seen:
            seen.append(url)
    return seen


def main():
    args = parse_args()
    conn = connect(args)
    cursor = conn.cursor(dictionary=True)

    where = "qr.qc_score IS NULL"
    params: tuple = ()
    if args.ids.strip():
        ids = [int(x) for x in args.ids.split(",") if x.strip()]
        where = f"qr.id IN ({','.join(['%s'] * len(ids))})"
        params = tuple(ids)

    cursor.execute(
        f"""
        SELECT qr.id, qr.tracker_id, qr.project_id, qr.status, qr.qc_status, qr.qc_score,
               qr.error_list, qr.qc_generated_count, qr.qc_file_path, qr.updated_at,
               u.user_name AS agent_name,
               (SELECT COUNT(*) FROM qc_correction_history ch WHERE ch.qc_record_id = qr.id) AS corrections,
               (SELECT COUNT(*) FROM qc_rework_history rh WHERE rh.qc_record_id = qr.id) AS reworks
        FROM qc_records qr
        LEFT JOIN tfs_user u ON u.user_id = qr.agent_id
        WHERE {where}
        ORDER BY qr.id
        """,
        params,
    )
    records = cursor.fetchall()
    print(f"Records to check: {len(records)}  ({'APPLY' if args.apply else 'DRY RUN'})\n")

    afd_cache = {}
    report_rows = []
    updates = []

    for rec in records:
        result = {
            "qc_record_id": rec["id"],
            "tracker_id": rec["tracker_id"],
            "agent": rec["agent_name"],
            "status": rec["status"],
            "corrections": rec["corrections"],
            "reworks": rec["reworks"],
            "qc_generated_count": rec["qc_generated_count"],
            "sample_rows": "",
            "errors": "",
            "new_score": "",
            "note": "",
            "file": "",
        }
        if rec["project_id"] not in afd_cache:
            afd_cache[rec["project_id"]] = load_afd(cursor, rec["project_id"])
        afd = afd_cache[rec["project_id"]]

        urls = sample_urls(cursor, rec)
        if not urls:
            result["note"] = "no saved QC sample file - skipped"
            report_rows.append(result)
            print(f"[skip] QC {rec['id']} tracker {rec['tracker_id']}: no sample file")
            continue

        parsed = None
        no_error_column = False
        for url in urls:
            try:
                parsed = read_sample(download_workbook(url), afd)
                result["file"] = url
                break
            except NoErrorColumn:
                no_error_column = True
                result["file"] = url
            except Exception as exc:  # noqa: BLE001
                result["note"] = f"could not read {url}: {exc}"
        if not parsed:
            if no_error_column:
                result["note"] = "sample saved without errors in the file - score cannot be recovered"
            report_rows.append(result)
            print(f"[skip] QC {rec['id']} tracker {rec['tracker_id']} {rec['agent_name']}: {result['note']}")
            continue

        row_count, error_list, per_row, summary = parsed
        if row_count == 0:
            result["note"] = "sample file has no data rows - skipped"
            report_rows.append(result)
            print(f"[skip] QC {rec['id']} tracker {rec['tracker_id']}: empty sample")
            continue

        notes = []
        scores = [record_score(row_errors, afd) for row_errors in per_row]
        new_score = round(sum(scores) / len(scores) + 1e-9, 2)
        if summary is not None:
            summary_total = sum(count for _, count in summary)
            if summary_total != len(error_list):
                summary_score, fatal = score_from_summary(summary, row_count, afd)
                if fatal:
                    result["note"] = (
                        f"rows show {len(error_list)} errors but Error List shows {summary_total} "
                        "with a fatal error - check manually"
                    )
                    report_rows.append(result)
                    print(f"[skip] QC {rec['id']} tracker {rec['tracker_id']}: {result['note']}")
                    continue
                new_score = round(summary_score + 1e-9, 2)
                notes.append(f"rows show {len(error_list)} errors, Error List shows {summary_total}; score uses Error List")
        if afd is None:
            notes.append("no error points set up for this project - every record counted as 100")
        if rec["qc_generated_count"] and int(rec["qc_generated_count"]) != row_count:
            notes.append(f"sample has {row_count} rows, record says {rec['qc_generated_count']}")

        result.update(
            {
                "sample_rows": row_count,
                "errors": len(error_list),
                "new_score": new_score,
                "note": "; ".join(notes),
            }
        )
        report_rows.append(result)
        updates.append((new_score, json.dumps(error_list), rec["id"]))
        print(
            f"[ok]   QC {rec['id']} tracker {rec['tracker_id']} {rec['agent_name']}: "
            f"{row_count} rows, {len(error_list)} errors -> {new_score}%"
            + (f"  ({result['note']})" if result["note"] else "")
        )

    report_path = args.report or os.path.join(
        os.path.dirname(__file__), f"repair_qc_null_scores_{datetime.now():%Y%m%d_%H%M%S}.csv"
    )
    with open(report_path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(report_rows[0].keys()) if report_rows else ["qc_record_id"])
        writer.writeheader()
        writer.writerows(report_rows)
    print(f"\nReport: {report_path}")
    print(f"Recalculated: {len(updates)}   Skipped: {len(report_rows) - len(updates)}")

    if not args.apply:
        print("\n[DRY RUN] Nothing was saved. Run again with --apply to write these scores.")
        cursor.close()
        conn.close()
        return

    try:
        if conn.in_transaction:
            conn.rollback()
        conn.start_transaction()
        for new_score, error_json, qc_id in updates:
            cursor.execute(
                "UPDATE qc_records SET qc_score = %s, error_list = %s, updated_at = updated_at "
                "WHERE id = %s AND qc_score IS NULL",
                (new_score, error_json, qc_id),
            )
        conn.commit()
        print(f"Saved {len(updates)} record(s).")
    except Exception:
        conn.rollback()
        raise
    finally:
        cursor.close()
        conn.close()


if __name__ == "__main__":
    main()
