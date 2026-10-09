from flask import Blueprint, request, jsonify, send_file
from config import get_db_connection
from datetime import datetime
from utils.cloudinary_utils import upload_to_cloudinary, delete_from_cloudinary
from io import BytesIO
import json
import urllib.request

qc_audit_bp = Blueprint("qc_audit", __name__)

FOLDER_QC_AUDIT = "hrms/qc_audit_files"

@qc_audit_bp.route("/add", methods=["POST"])
def create_qc_audit():

    form = request.form

    qc_record_id = form.get("qc_record_id")
    qc_score = form.get("qc_score")
    error_notes = form.get("error_notes")

    if not qc_record_id or not qc_score:
        return jsonify({
            "status":400,
            "message":"qc_record_id and qc_score required"
        }),400

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    try:

        qc_file_url = None
        uploaded = request.files.get("qc_checked_file")

        if uploaded and uploaded.filename:
            extension = uploaded.filename.rsplit('.', 1)[1].lower()
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

            custom_name = f"qc_checked_file_{qc_record_id}_{timestamp}.{extension}"

            qc_file_url, _ = upload_to_cloudinary(
                uploaded,
                FOLDER_QC_AUDIT,
                display_name=custom_name ,
                resource_type="raw"
            )

        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        cursor.execute("""
        INSERT INTO qc_audit
        (qc_record_id,qc_score,qc_checked_file,error_notes,created_date,updated_date)
        VALUES(%s,%s,%s,%s,%s,%s)
        """,(
            qc_record_id,
            qc_score,
            qc_file_url,
            error_notes,
            now,
            now
        ))

        conn.commit()

        return jsonify({
            "status":201,
            "message":"QC audit created"
        }),201

    except Exception as e:
        conn.rollback()
        return jsonify({
            "status":500,
            "message":str(e)
        }),500

    finally:
        cursor.close()
        conn.close() 
        

@qc_audit_bp.route("/test", methods=["GET"])
def test_qc_audit():
    return jsonify({
        "status": 200,
        "message": "QC Audit Test - Code Updated Successfully",
        "timestamp": str(datetime.now())
    })

def _valid_iso_date(value):
    text = (value or "").strip()
    if not text:
        return None
    try:
        datetime.strptime(text[:10], "%Y-%m-%d")
        return text[:10]
    except ValueError:
        return None


@qc_audit_bp.route("/report", methods=["POST"])
def qc_audit_report():

    data = request.get_json() or {}
    start_date = _valid_iso_date(data.get("start_date"))
    end_date = _valid_iso_date(data.get("end_date"))
    qc_start_date = _valid_iso_date(data.get("qc_start_date"))
    qc_end_date = _valid_iso_date(data.get("qc_end_date"))

    # If no date filters provided, default to current month on worked date.
    if not any((start_date, end_date, qc_start_date, qc_end_date)):
        from calendar import monthrange
        today = datetime.now()
        start_date = today.replace(day=1).strftime("%Y-%m-%d")
        last_day = monthrange(today.year, today.month)[1]
        end_date = today.replace(day=last_day).strftime("%Y-%m-%d")
    
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)

    try:

        query = """
        SELECT
        qa.created_date AS audit_datetime,
        qr.date_of_file_submission AS worked_date,
        qr.created_at AS qc_date,
        qr.updated_at AS evaluation_date,
        tu.user_name AS agent_name,
        qc_user.user_name AS qc_agent_name,
        p.project_name AS project,
        t.task_name AS task,
        qr.qc_generated_count AS total_qcs,
        qa.qc_score AS avg_qc_score,
        qr.error_list AS total_errors,
        qa.qc_checked_file,
        qr.qc_status,
        qa.error_notes

        FROM qc_audit qa

        LEFT JOIN qc_records qr
        ON qa.qc_record_id = qr.id

        LEFT JOIN tfs_user tu
        ON qr.agent_id = tu.user_id

        LEFT JOIN tfs_user qc_user
        ON qr.qa_user_id = qc_user.user_id

        LEFT JOIN project p
        ON qr.project_id = p.project_id

        LEFT JOIN task t
        ON qr.task_id = t.task_id
        """

        where = []
        params = []
        if start_date:
            where.append("DATE(qr.date_of_file_submission) >= %s")
            params.append(start_date)
        if end_date:
            where.append("DATE(qr.date_of_file_submission) <= %s")
            params.append(end_date)
        if qc_start_date:
            where.append("DATE(qr.created_at) >= %s")
            params.append(qc_start_date)
        if qc_end_date:
            where.append("DATE(qr.created_at) <= %s")
            params.append(qc_end_date)
        if where:
            query += " WHERE " + " AND ".join(where)

        query += " ORDER BY qa.created_date DESC"

        cursor.execute(query, params)
        rows = cursor.fetchall()

        return jsonify({
            "status": 200,
            "message": "QC Audit Report",
            "data": {
                "count": len(rows),
                "records": rows
            }
        }), 200

    except Exception as e:
        return jsonify({
            "status": 500,
            "message": str(e)
        }), 500

    finally:
        cursor.close()
        conn.close()


def _parse_error_list(raw):
    if not raw:
        return []
    if isinstance(raw, list):
        return raw
    if isinstance(raw, (bytes, bytearray)):
        raw = raw.decode("utf-8", errors="ignore")
    if isinstance(raw, str):
        try:
            parsed = json.loads(raw)
            return parsed if isinstance(parsed, list) else []
        except Exception:
            return []
    return []


def _error_label(err):
    if err is None:
        return ""
    if not isinstance(err, dict):
        return str(err)
    return (
        err.get("error")
        or (
            f"{err.get('category')} - {err.get('subcategory')}"
            if err.get("category") and err.get("subcategory")
            else ""
        )
        or err.get("name")
        or err.get("message")
        or json.dumps(err)
    )


def _annotate_openpyxl_sheet(ws, error_list):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    errors = _parse_error_list(error_list)
    last_col = ws.max_column or 1
    error_col = last_col
    header = ws.cell(1, last_col).value
    if str(header or "").strip().lower() != "errors":
        error_col = last_col + 1

    header_cell = ws.cell(1, error_col, "Errors")
    header_cell.font = Font(bold=True, color="FFFFFF")
    header_cell.fill = PatternFill("solid", fgColor="B91C1C")
    header_cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions[get_column_letter(error_col)].width = 48

    by_row = {}
    unique = {}
    for err in errors:
        label = (_error_label(err) or "").strip()
        if not label:
            continue
        unique[label] = unique.get(label, 0) + 1
        try:
            row_num = int(err.get("row")) if isinstance(err, dict) else 0
        except (TypeError, ValueError):
            continue
        if row_num < 1:
            continue
        by_row.setdefault(row_num, [])
        if label not in by_row[row_num]:
            by_row[row_num].append(label)

    last_data = 1
    for r in range(2, (ws.max_row or 1) + 1):
        val = str(ws.cell(r, error_col).value or "").strip().lower()
        if val == "error list":
            break
        last_data = r

    pink = PatternFill("solid", fgColor="FFC7CE")
    light = PatternFill("solid", fgColor="FFEBEE")
    red_font = Font(bold=True, color="9C0006")
    for excel_row, labels in by_row.items():
        cell = ws.cell(excel_row, error_col, "; ".join(labels))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = red_font
        cell.fill = pink
        for c in range(1, error_col):
            data_cell = ws.cell(excel_row, c)
            if data_cell.fill is None or data_cell.fill.fill_type is None:
                data_cell.fill = light
        if excel_row > last_data:
            last_data = excel_row

    next_row = last_data + 2
    title = ws.cell(next_row, error_col, "Error List")
    title.font = Font(bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="B91C1C")
    next_row += 1
    if not unique:
        ws.cell(next_row, error_col, "No errors")
        return
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for name, count in unique.items():
        cell = ws.cell(next_row, error_col, f"{name} ({count})")
        cell.font = Font(color="9C0006")
        cell.fill = yellow
        next_row += 1


@qc_audit_bp.route("/download_annotated/<int:qc_id>", methods=["GET"])
def download_annotated_qc_file(qc_id):
    """Build an Excel with Errors column + highlighted rows for a QC record."""
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute(
            "SELECT id, qc_file_path, error_list FROM qc_records WHERE id = %s LIMIT 1",
            (qc_id,),
        )
        record = cursor.fetchone()
        if not record:
            return jsonify({"status": 404, "message": "QC record not found"}), 404

        from openpyxl import Workbook, load_workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "QC Sample"
        file_url = (record.get("qc_file_path") or "").strip()
        if len(file_url) > 1 and "https://" in file_url[1:]:
            file_url = file_url[file_url.find("https://"):]
        elif len(file_url) > 1 and "http://" in file_url[1:]:
            file_url = file_url[file_url.find("http://"):]

        if file_url.startswith("http"):
            try:
                req = urllib.request.Request(
                    file_url,
                    headers={"User-Agent": "HRMS-QC-Annotate/1.0"},
                )
                with urllib.request.urlopen(req, timeout=60) as resp:
                    payload = BytesIO(resp.read())
                wb = load_workbook(payload)
                ws = wb.active
            except Exception as load_err:
                print("annotated download source load failed:", load_err)
                ws["A1"] = "Agent file could not be loaded. Error list is below."

        _annotate_openpyxl_sheet(ws, record.get("error_list"))
        out = BytesIO()
        wb.save(out)
        out.seek(0)
        return send_file(
            out,
            as_attachment=True,
            download_name=f"QC_Errors_Record_{qc_id}.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    except Exception as e:
        return jsonify({"status": 500, "message": str(e)}), 500
    finally:
        cursor.close()
        conn.close()