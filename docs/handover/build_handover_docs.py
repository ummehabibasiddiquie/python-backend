"""
Build the developer handover documents (Word) from handover_content*.py.

Usage (from backend-api):
    python docs/handover/build_handover_docs.py
Output: docs/handover/HRMS Developer Guide/<NN> - <title>.docx
"""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "sop"))
sys.path.insert(0, str(HERE))

from build_sop_docs import BRAND, Renderer, set_cell_shading  # noqa: E402

OUT = HERE / "HRMS Developer Guide"
VERSION = "1.0"
ISSUED = date(2026, 10, 22).strftime("%d %b %Y")


class HandoverRenderer(Renderer):
    def code(self, text: str) -> None:
        table = self.doc.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        cell = table.rows[0].cells[0]
        set_cell_shading(cell, "F1F5F9")
        paragraph = cell.paragraphs[0]
        for index, line in enumerate(text.strip("\n").split("\n")):
            if index:
                paragraph.add_run().add_break()
            run = paragraph.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        self.doc.add_paragraph()

    def render(self, blocks: list) -> None:
        for block in blocks:
            if block[0] == "code":
                self.code(block[1])
            else:
                super().render([block])


def build(spec: dict, path: Path | None = None) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for section in doc.sections:
        if spec.get("landscape"):
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width
        section.left_margin = section.right_margin = Inches(0.8)
        section.top_margin = section.bottom_margin = Inches(0.8)
        footer = section.footer.paragraphs[0]
        footer.text = f"TFS HRMS - Developer Handover - {spec['title']} - v{VERSION}"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title = doc.add_heading(f"Handover {spec['number']}: {spec['title']}", level=0)
    for run in title.runs:
        run.font.color.rgb = BRAND

    meta = doc.add_table(rows=4, cols=2)
    meta.style = "Table Grid"
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row, (label, value) in zip(meta.rows, [
        ("Document", f"{spec['number']} - {spec['title']}"),
        ("For", spec["audience"]),
        ("Version", VERSION),
        ("Issued", ISSUED),
    ]):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], "DBEAFE")
        row.cells[0].paragraphs[0].runs[0].bold = True
    doc.add_paragraph()

    HandoverRenderer(doc).render(spec["blocks"])

    if path is None:
        OUT.mkdir(parents=True, exist_ok=True)
        path = OUT / f"{spec['number']} - {spec['title']}.docx"
    doc.save(str(path))
    return path


if __name__ == "__main__":
    from handover_content import HANDOVER_DOCS

    problems = []
    written = set()
    for spec in HANDOVER_DOCS:
        try:
            path = build(spec)
        except PermissionError as exc:
            problems.append(f"could not write {spec['number']} - {spec['title']}.docx (open in Word?): {exc}")
            written.add(f"{spec['number']} - {spec['title']}.docx")
            continue
        written.add(path.name)
        print("wrote", path.name)

    # Remove .docx files left over from earlier runs (old numbering / titles).
    for old in sorted(OUT.glob("*.docx")):
        if old.name in written or old.name.startswith("~$"):
            continue
        try:
            old.unlink()
            print("removed old", old.name)
        except PermissionError as exc:
            problems.append(f"could not delete old file {old.name} (open in Word?): {exc}")

    if problems:
        print("\nPROBLEMS:")
        for problem in problems:
            print(" -", problem)
        sys.exit(1)
